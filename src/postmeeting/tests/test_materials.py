"""付帯資料の取込（materials.py）のテスト（③付帯資料）。"""

from __future__ import annotations

from pathlib import Path

import pytest

from subtext_postmeeting import materials
from subtext_postmeeting.errors import PipelineError


def _write(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


# --- extract_one: .txt / .md -------------------------------------------------


def test_extract_one_reads_txt_as_utf8(tmp_path: Path) -> None:
    path = _write(tmp_path / "note.txt", "会議の要点メモ")
    result = materials.extract_one(path)

    assert result.file_name == "note.txt"
    assert result.kind == "txt"
    assert result.extraction_method == "text"
    assert result.text == "会議の要点メモ"
    assert result.char_count == len("会議の要点メモ")


def test_extract_one_reads_md(tmp_path: Path) -> None:
    path = _write(tmp_path / "agenda.md", "# 議題\n- A\n- B\n")
    result = materials.extract_one(path)
    assert result.kind == "md"
    assert "# 議題" in result.text


def test_extract_one_falls_back_to_cp932(tmp_path: Path) -> None:
    path = tmp_path / "legacy.txt"
    path.write_bytes("シフトJISの資料".encode("cp932"))
    result = materials.extract_one(path)
    assert result.text == "シフトJISの資料"


def test_extract_one_raises_when_encoding_unrecognized(tmp_path: Path) -> None:
    path = tmp_path / "binary.txt"
    path.write_bytes(b"\x81\xff\x00\x01\x02\x03")  # UTF-8 でも CP932 でも不正なバイト列
    with pytest.raises(PipelineError, match="文字コード"):
        materials.extract_one(path)


def test_extract_one_rejects_unsupported_extension(tmp_path: Path) -> None:
    path = _write(tmp_path / "image.png", "not really an image")
    with pytest.raises(PipelineError, match="未対応"):
        materials.extract_one(path)


# --- extract_one: .pdf --------------------------------------------------------


def _write_blank_pdf(path: Path) -> Path:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as fh:
        writer.write(fh)
    return path


def test_extract_one_rejects_pdf_with_empty_text_layer(tmp_path: Path) -> None:
    path = _write_blank_pdf(tmp_path / "scanned.pdf")
    with pytest.raises(PipelineError, match="テキスト層が空"):
        materials.extract_one(path)


def test_extract_one_raises_actionable_error_for_corrupt_pdf(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(PipelineError, match="PDF の読込に失敗"):
        materials.extract_one(path)


def test_extract_one_pdf_with_text_layer(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """実際のテキスト層抽出は pypdf.PdfReader をフェイクに差し替えて検証する（純粋な組立の確認）。"""
    path = _write_blank_pdf(tmp_path / "doc.pdf")

    class _FakePage:
        def extract_text(self) -> str:
            return "資料の本文"

    class _FakeReader:
        def __init__(self, _path: str) -> None:
            self.pages = [_FakePage(), _FakePage()]

    monkeypatch.setattr("pypdf.PdfReader", _FakeReader)

    result = materials.extract_one(path)

    assert result.kind == "pdf"
    assert result.extraction_method == "pdf-text-layer"
    assert result.text == "資料の本文\n資料の本文"


# --- extract_one: .pptx -------------------------------------------------------


def _write_pptx(path: Path, *, title: str, body: str, notes: str = "") -> Path:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))
    return path


def test_extract_one_pptx_reads_slide_text_and_notes(tmp_path: Path) -> None:
    path = _write_pptx(tmp_path / "slides.pptx", title="提案概要", body="ポイント1", notes="発表者メモ")

    result = materials.extract_one(path)

    assert result.kind == "pptx"
    assert result.extraction_method == "pptx"
    assert "提案概要" in result.text
    assert "ポイント1" in result.text
    assert "発表者メモ" in result.text
    assert "[スライド1]" in result.text


def test_extract_one_pptx_without_notes_omits_notes_marker(tmp_path: Path) -> None:
    path = _write_pptx(tmp_path / "slides.pptx", title="T", body="B")
    result = materials.extract_one(path)
    assert "[ノート]" not in result.text


def test_extract_one_raises_actionable_error_for_corrupt_pptx(tmp_path: Path) -> None:
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"not a pptx")
    with pytest.raises(PipelineError, match="pptx の読込に失敗"):
        materials.extract_one(path)


# --- content_hash -------------------------------------------------------------


def test_content_hash_changes_with_content(tmp_path: Path) -> None:
    a = _write(tmp_path / "a.txt", "A")
    b = _write(tmp_path / "b.txt", "B")
    assert materials.content_hash(a) != materials.content_hash(b)


def test_content_hash_stable_for_same_content(tmp_path: Path) -> None:
    a = _write(tmp_path / "a.txt", "同じ内容")
    b = _write(tmp_path / "b.txt", "同じ内容")
    assert materials.content_hash(a) == materials.content_hash(b)


# --- collect -------------------------------------------------------------------


def test_collect_includes_all_when_under_limit(tmp_path: Path) -> None:
    a = _write(tmp_path / "a.txt", "12345")
    b = _write(tmp_path / "b.txt", "67890")
    result = materials.collect((a, b), max_total_chars=1000)
    assert [m.file_name for m in result.included] == ["a.txt", "b.txt"]
    assert result.excluded == ()
    assert result.warnings == ()


def test_collect_excludes_files_that_exceed_total_limit_in_order(tmp_path: Path) -> None:
    a = _write(tmp_path / "a.txt", "12345")  # 5文字
    b = _write(tmp_path / "b.txt", "1234567890")  # 10文字
    result = materials.collect((a, b), max_total_chars=8)
    assert [m.file_name for m in result.included] == ["a.txt"]
    assert result.excluded == (("b.txt", "文字数上限（8）を超えるため除外"),)


def test_collect_does_not_split_a_single_file_at_the_boundary(tmp_path: Path) -> None:
    """1ファイルの途中で切らない。丁度収まらないファイルは丸ごと除外する。"""
    a = _write(tmp_path / "a.txt", "1234567890")  # 10文字
    result = materials.collect((a,), max_total_chars=5)
    assert result.included == ()
    assert result.excluded == (("a.txt", "文字数上限（5）を超えるため除外"),)


def test_collect_converts_extraction_failures_to_warnings_not_exceptions(tmp_path: Path) -> None:
    good = _write(tmp_path / "good.txt", "OK")
    bad = _write(tmp_path / "bad.png", "not an image")
    result = materials.collect((good, bad), max_total_chars=1000)
    assert [m.file_name for m in result.included] == ["good.txt"]
    assert len(result.warnings) == 1
    assert "bad.png" in result.warnings[0]


def test_collect_empty_input_returns_empty_result() -> None:
    result = materials.collect((), max_total_chars=1000)
    assert result.is_empty()


# --- format_for_prompt --------------------------------------------------------


def test_format_for_prompt_empty_when_no_materials() -> None:
    assert materials.format_for_prompt(materials.MaterialsResult()) == ""


def test_format_for_prompt_includes_filename_and_char_count(tmp_path: Path) -> None:
    a = _write(tmp_path / "a.txt", "本文")
    result = materials.collect((a,), max_total_chars=1000)
    prompt = materials.format_for_prompt(result)
    assert prompt.startswith("[参考資料]")
    assert "資料1: a.txt（2文字）" in prompt
    assert "本文" in prompt


def test_format_for_prompt_numbers_multiple_materials(tmp_path: Path) -> None:
    a = _write(tmp_path / "a.txt", "A")
    b = _write(tmp_path / "b.txt", "B")
    result = materials.collect((a, b), max_total_chars=1000)
    prompt = materials.format_for_prompt(result)
    assert "資料1: a.txt" in prompt
    assert "資料2: b.txt" in prompt


# --- MaterialsResult / MaterialText の JSON round-trip ------------------------


def test_materials_result_to_json_round_trip(tmp_path: Path) -> None:
    a = _write(tmp_path / "a.txt", "本文")
    b = _write(tmp_path / "b.png", "ignored")
    result = materials.collect((a, b), max_total_chars=1000)

    restored = materials.MaterialsResult.from_json(result.to_json())

    assert restored == result


# --- キャッシュ（load_cache/save_cache/needs_extraction） ---------------------


def test_load_cache_missing_file_returns_none(tmp_path: Path) -> None:
    assert materials.load_cache(tmp_path / "materials.extracted.json") is None


def test_load_cache_raises_actionable_error_on_broken_json(tmp_path: Path) -> None:
    path = tmp_path / "materials.extracted.json"
    path.write_text("{not json", encoding="utf-8")
    with pytest.raises(PipelineError, match="materials.extracted.json"):
        materials.load_cache(path)


def test_save_and_load_cache_round_trip(tmp_path: Path) -> None:
    a = _write(tmp_path / "materials" / "a.txt", "本文")
    result = materials.collect((a,), max_total_chars=1000)
    cache_path = tmp_path / "materials.extracted.json"

    materials.save_cache(cache_path, result)

    assert materials.load_cache(cache_path) == result


def test_needs_extraction_true_when_no_cache(tmp_path: Path) -> None:
    materials_dir = tmp_path / "materials"
    materials_dir.mkdir()
    assert materials.needs_extraction(materials_dir, None) is True


def test_needs_extraction_false_when_unchanged(tmp_path: Path) -> None:
    materials_dir = tmp_path / "materials"
    materials_dir.mkdir()
    a = _write(materials_dir / "a.txt", "本文")
    cache = materials.collect((a,), max_total_chars=1000)
    assert materials.needs_extraction(materials_dir, cache) is False


def test_needs_extraction_true_when_file_added(tmp_path: Path) -> None:
    materials_dir = tmp_path / "materials"
    materials_dir.mkdir()
    a = _write(materials_dir / "a.txt", "本文")
    cache = materials.collect((a,), max_total_chars=1000)
    _write(materials_dir / "b.txt", "追加資料")
    assert materials.needs_extraction(materials_dir, cache) is True


def test_needs_extraction_true_when_file_removed(tmp_path: Path) -> None:
    materials_dir = tmp_path / "materials"
    materials_dir.mkdir()
    a = _write(materials_dir / "a.txt", "本文")
    cache = materials.collect((a,), max_total_chars=1000)
    a.unlink()
    assert materials.needs_extraction(materials_dir, cache) is True


def test_needs_extraction_true_when_file_content_changed(tmp_path: Path) -> None:
    materials_dir = tmp_path / "materials"
    materials_dir.mkdir()
    a = _write(materials_dir / "a.txt", "本文")
    cache = materials.collect((a,), max_total_chars=1000)
    _write(a, "変更後の本文")
    assert materials.needs_extraction(materials_dir, cache) is True


def test_needs_extraction_false_when_dir_absent_and_cache_empty() -> None:
    empty_cache = materials.MaterialsResult()
    assert materials.needs_extraction(Path("/does/not/exist"), empty_cache) is False


def test_needs_extraction_true_when_dir_removed_but_cache_had_materials(tmp_path: Path) -> None:
    materials_dir = tmp_path / "materials"
    materials_dir.mkdir()
    a = _write(materials_dir / "a.txt", "本文")
    cache = materials.collect((a,), max_total_chars=1000)
    a.unlink()
    materials_dir.rmdir()
    assert materials.needs_extraction(materials_dir, cache) is True
